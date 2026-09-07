from flask import Flask, render_template, request, send_from_directory
from groq import Groq
from werkzeug.utils import secure_filename

import os
import io
import re
import tempfile

import pandas as pd
import fitz
import pytesseract

from PIL import Image, ImageOps, ImageFilter
from docx import Document
from dotenv import load_dotenv


# =========================================================
# TESSERACT
# =========================================================

TESSERACT_PATH = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

pytesseract.pytesseract.tesseract_cmd = TESSERACT_PATH


# =========================================================
# ENVIRONMENT
# =========================================================

load_dotenv()

app = Flask(__name__)

app.config["MAX_CONTENT_LENGTH"] = 50 * 1024 * 1024


# =========================================================
# GROQ
# =========================================================

api_key = os.getenv("GROQ_API_KEY")

if not api_key:
    raise ValueError(
        "GROQ_API_KEY is missing. "
        "Please add GROQ_API_KEY to your .env file."
    )

client = Groq(api_key=api_key)

MODEL = "openai/gpt-oss-120b"


# =========================================================
# LANGUAGES
# =========================================================

LANGUAGES = [
    "Abkhaz",
    "Acehnese",
    "Acholi",
    "Afar",
    "Afrikaans",
    "Albanian",
    "Alur",
    "Amharic",
    "Arabic",
    "Armenian",
    "Assamese",
    "Avar",
    "Awadhi",
    "Aymara",
    "Azerbaijani",
    "Balinese",
    "Baluchi",
    "Bambara",
    "Baoulé",
    "Bashkir",
    "Basque",
    "Batak Karo",
    "Batak Simalungun",
    "Batak Toba",
    "Belarusian",
    "Bemba",
    "Bengali",
    "Betawi",
    "Bhojpuri",
    "Bikol",
    "Bosnian",
    "Breton",
    "Bulgarian",
    "Buryat",
    "Cantonese",
    "Catalan",
    "Cebuano",
    "Chamorro",
    "Chechen",
    "Chichewa",
    "Chinese (Simplified)",
    "Chinese (Traditional)",
    "Chuukese",
    "Chuvash",
    "Corsican",
    "Crimean Tatar (Cyrillic)",
    "Crimean Tatar (Latin)",
    "Croatian",
    "Czech",
    "Danish",
    "Dari",
    "Dhivehi",
    "Dinka",
    "Dogri",
    "Dombe",
    "Dutch",
    "Dyula",
    "Dzongkha",
    "English",
    "Esperanto",
    "Estonian",
    "Ewe",
    "Faroese",
    "Fijian",
    "Filipino",
    "Finnish",
    "Fon",
    "French",
    "French (Canada)",
    "Frisian",
    "Friulian",
    "Fulani",
    "Ga",
    "Galician",
    "Georgian",
    "German",
    "Greek",
    "Guarani",
    "Gujarati",
    "Haitian Creole",
    "Hakha Chin",
    "Hausa",
    "Hawaiian",
    "Hebrew",
    "Hiligaynon",
    "Hindi",
    "Hmong",
    "Hungarian",
    "Hunsrik",
    "Iban",
    "Icelandic",
    "Igbo",
    "Ilocano",
    "Indonesian",
    "Inuktut (Latin)",
    "Inuktut (Syllabics)",
    "Irish",
    "Italian",
    "Jamaican Patois",
    "Japanese",
    "Javanese",
    "Jingpo",
    "Kalaallisut",
    "Kannada",
    "Kanuri",
    "Kapampangan",
    "Kazakh",
    "Khasi",
    "Khmer",
    "Kiga",
    "Kikongo",
    "Kinyarwanda",
    "Kituba",
    "Kokborok",
    "Komi",
    "Konkani",
    "Korean",
    "Krio",
    "Kurdish (Kurmanji)",
    "Kurdish (Sorani)",
    "Kyrgyz",
    "Lao",
    "Latgalian",
    "Latin",
    "Latvian",
    "Ligurian",
    "Limburgish",
    "Lingala",
    "Lithuanian",
    "Lombard",
    "Luganda",
    "Luo",
    "Luxembourgish",
    "Macedonian",
    "Madurese",
    "Maithili",
    "Makassar",
    "Malagasy",
    "Malay",
    "Malay (Jawi)",
    "Malayalam",
    "Maltese",
    "Mam",
    "Manx",
    "Maori",
    "Marathi",
    "Marshallese",
    "Marwadi",
    "Mauritian Creole",
    "Meadow Mari",
    "Meiteilon (Manipuri)",
    "Minang",
    "Mizo",
    "Mongolian",
    "Myanmar (Burmese)",
    "Nahuatl (Eastern Huasteca)",
    "Ndau",
    "Ndebele (South)",
    "Nepalbhasa (Newari)",
    "Nepali",
    "NKo",
    "Norwegian (Bokmål)",
    "Nuer",
    "Occitan",
    "Odia (Oriya)",
    "Oromo",
    "Ossetian",
    "Pangasinan",
    "Papiamento",
    "Pashto",
    "Persian",
    "Polish",
    "Portuguese (Brazil)",
    "Portuguese (Portugal)",
    "Punjabi (Gurmukhi)",
    "Punjabi (Shahmukhi)",
    "Quechua",
    "Q'eqchi'",
    "Romani",
    "Romanian",
    "Rundi",
    "Russian",
    "Sami (North)",
    "Samoan",
    "Sango",
    "Sanskrit",
    "Santali (Latin)",
    "Santali (Ol Chiki)",
    "Scots Gaelic",
    "Sepedi",
    "Serbian",
    "Sesotho",
    "Seychellois Creole",
    "Shan",
    "Shona",
    "Sicilian",
    "Silesian",
    "Sindhi",
    "Sinhala",
    "Slovak",
    "Slovenian",
    "Somali",
    "Spanish",
    "Sundanese",
    "Susu",
    "Swahili",
    "Swati",
    "Swedish",
    "Tahitian",
    "Tajik",
    "Tamazight",
    "Tamazight (Tifinagh)",
    "Tamil",
    "Tatar",
    "Telugu",
    "Tetum",
    "Thai",
    "Tibetan",
    "Tigrinya",
    "Tiv",
    "Tok Pisin",
    "Tongan",
    "Tshiluba",
    "Tsonga",
    "Tswana",
    "Tulu",
    "Tumbuka",
    "Turkish",
    "Turkmen",
    "Tuvan",
    "Twi",
    "Udmurt",
    "Ukrainian",
    "Urdu",
    "Uyghur",
    "Uzbek",
    "Venda",
    "Venetian",
    "Vietnamese",
    "Waray",
    "Welsh",
    "Wolof",
    "Xhosa",
    "Yakut",
    "Yiddish",
    "Yoruba",
    "Yucatec Maya",
    "Zapotec",
    "Zulu"
]


# =========================================================
# TESSERACT LANGUAGE CODES
# =========================================================

TESSERACT_LANG_CODES = {

    "English": "eng",
    "Tamil": "tam",
    "Hindi": "hin",
    "Telugu": "tel",
    "Malayalam": "mal",
    "Kannada": "kan",
    "Bengali": "ben",
    "Gujarati": "guj",
    "Marathi": "mar",

    "Punjabi (Gurmukhi)": "pan",
    "Punjabi (Shahmukhi)": "urd",
    "Urdu": "urd",

    "Arabic": "ara",
    "Persian": "fas",

    "French": "fra",
    "French (Canada)": "fra",

    "German": "deu",
    "Spanish": "spa",
    "Italian": "ita",

    "Portuguese (Brazil)": "por",
    "Portuguese (Portugal)": "por",

    "Russian": "rus",
    "Ukrainian": "ukr",
    "Greek": "ell",
    "Hebrew": "heb",

    "Japanese": "jpn",
    "Korean": "kor",

    "Chinese (Simplified)": "chi_sim",
    "Chinese (Traditional)": "chi_tra",

    "Thai": "tha",
    "Vietnamese": "vie",
    "Indonesian": "ind",
    "Malay": "msa",
    "Turkish": "tur",
    "Dutch": "nld",
    "Polish": "pol",
    "Romanian": "ron",
    "Swedish": "swe",
    "Danish": "dan",
    "Finnish": "fin",
    "Norwegian (Bokmål)": "nor",
    "Czech": "ces",
    "Hungarian": "hun",
    "Slovak": "slk",
    "Slovenian": "slv",
    "Croatian": "hrv",
    "Serbian": "srp",
    "Bulgarian": "bul",
    "Estonian": "est",
    "Latvian": "lav",
    "Lithuanian": "lit",
    "Nepali": "nep",
    "Sinhala": "sin",
    "Swahili": "swa",
    "Afrikaans": "afr",

    "Odia (Oriya)": "ori",
    "Assamese": "asm",
    "Sanskrit": "san",
    "Sindhi": "snd",
    "Konkani": "kok",
    "Maithili": "mai",
    "Dogri": "doi",
    "Nepalbhasa (Newari)": "new",
    "Marwadi": "mar"
}


# =========================================================
# GET INSTALLED TESSERACT LANGUAGES
# =========================================================

def get_installed_tesseract_languages():

    try:

        languages = pytesseract.get_languages(
            config=""
        )

        print(
            "Installed Tesseract languages:",
            languages
        )

        return languages

    except Exception as e:

        print(
            "Could not read Tesseract languages:",
            e
        )

        return []


# =========================================================
# OCR LANGUAGE
# =========================================================

def get_ocr_language(
    source_language="Auto Detect"
):

    installed = get_installed_tesseract_languages()

    if not installed:

        raise ValueError(
            "Tesseract is installed, but no OCR language "
            "data was found."
        )


    # -----------------------------------------------------
    # AUTO DETECT
    # -----------------------------------------------------

    if source_language == "Auto Detect":

        # Hindi + English mixed document

        if (
            "hin" in installed
            and "eng" in installed
        ):

            print(
                "AUTO OCR: Using hin+eng"
            )

            return "hin+eng"


        if "hin" in installed:

            print(
                "AUTO OCR: Using hin"
            )

            return "hin"


        if "eng" in installed:

            print(
                "AUTO OCR: Using eng"
            )

            return "eng"


        raise ValueError(
            "Auto Detect could not find a suitable "
            "Tesseract OCR language."
        )


    # -----------------------------------------------------
    # MANUAL SOURCE
    # -----------------------------------------------------

    code = TESSERACT_LANG_CODES.get(
        source_language
    )

    if not code:

        raise ValueError(
            f"OCR is not configured for "
            f"'{source_language}'."
        )


    if code not in installed:

        raise ValueError(
            f"Tesseract language '{code}' is not installed."
        )


    # -----------------------------------------------------
    # MIXED SOURCE + ENGLISH
    # -----------------------------------------------------

    if (
        code != "eng"
        and "eng" in installed
    ):

        combined = f"{code}+eng"

        print(
            f"MANUAL OCR: Using {combined}"
        )

        return combined


    print(
        f"MANUAL OCR: Using {code}"
    )

    return code


# =========================================================
# IMAGE PREPROCESSING
# =========================================================

def preprocess_image(image):

    image = image.convert("RGB")

    gray = ImageOps.grayscale(
        image
    )

    gray = ImageOps.autocontrast(
        gray
    )

    gray = gray.filter(
        ImageFilter.SHARPEN
    )

    return gray


# =========================================================
# OCR IMAGE
# =========================================================

def ocr_image(
    image,
    source_language="Auto Detect"
):

    ocr_language = get_ocr_language(
        source_language
    )

    print(
        "OCR language:",
        ocr_language
    )


    # =====================================================
    # OCR PASS 1
    # =====================================================

    original_text = pytesseract.image_to_string(

        image,

        lang=ocr_language,

        config="--oem 3 --psm 6"

    )


    # =====================================================
    # OCR PASS 2
    # =====================================================

    processed_image = preprocess_image(
        image
    )

    processed_text = pytesseract.image_to_string(

        processed_image,

        lang=ocr_language,

        config="--oem 3 --psm 6"

    )


    original_clean = (
        original_text.strip()
    )

    processed_clean = (
        processed_text.strip()
    )


    # -----------------------------------------------------
    # Choose better OCR result
    # -----------------------------------------------------

    if not original_clean:

        selected_text = processed_clean

    elif not processed_clean:

        selected_text = original_clean

    else:

        if len(processed_clean) > len(original_clean):

            selected_text = processed_clean

        else:

            selected_text = original_clean


    print(
        "OCR result length:",
        len(selected_text)
    )

    print(
        "OCR preview:",
        selected_text[:500]
    )

    return (
        selected_text,
        ocr_language
    )


# =========================================================
# DETECT LANGUAGE
# =========================================================

def detect_translation_source(
    text,
    selected_source
):

    if selected_source != "Auto Detect":

        return selected_source


    if not text:

        return "Auto Detect"


    # -----------------------------------------------------
    # Hindi / Devanagari
    # -----------------------------------------------------

    has_devanagari = bool(
        re.search(
            r"[\u0900-\u097F]",
            text
        )
    )


    # -----------------------------------------------------
    # Tamil
    # -----------------------------------------------------

    has_tamil = bool(
        re.search(
            r"[\u0B80-\u0BFF]",
            text
        )
    )


    # -----------------------------------------------------
    # Telugu
    # -----------------------------------------------------

    has_telugu = bool(
        re.search(
            r"[\u0C00-\u0C7F]",
            text
        )
    )


    # -----------------------------------------------------
    # Malayalam
    # -----------------------------------------------------

    has_malayalam = bool(
        re.search(
            r"[\u0D00-\u0D7F]",
            text
        )
    )


    # -----------------------------------------------------
    # Kannada
    # -----------------------------------------------------

    has_kannada = bool(
        re.search(
            r"[\u0C80-\u0CFF]",
            text
        )
    )


    # -----------------------------------------------------
    # Latin
    # -----------------------------------------------------

    has_latin = bool(
        re.search(
            r"[A-Za-z]",
            text
        )
    )


    if has_devanagari and has_latin:

        return "Hindi and English"


    if has_devanagari:

        return "Hindi"


    if has_tamil and has_latin:

        return "Tamil and English"


    if has_tamil:

        return "Tamil"


    if has_telugu and has_latin:

        return "Telugu and English"


    if has_telugu:

        return "Telugu"


    if has_malayalam:

        return "Malayalam"


    if has_kannada:

        return "Kannada"


    if has_latin:

        return "English"


    return "Auto Detect"


# =========================================================
# TRANSLATE TEXT
# =========================================================

def translate_text(
    text,
    target_language,
    source_language="Auto Detect"
):

    if not text or not text.strip():

        return ""


    clean_text = text.strip()


    print("=" * 60)

    print(
        "TRANSLATION REQUEST"
    )

    print(
        "Source:",
        source_language
    )

    print(
        "Target:",
        target_language
    )

    print(
        "Text length:",
        len(clean_text)
    )

    print(
        "Text preview:",
        clean_text[:300]
    )

    print("=" * 60)


    prompt = f"""
You are a professional document translator.

Translate the following document text from:
{source_language}

into:
{target_language}

IMPORTANT INSTRUCTIONS:

- Translate the content faithfully.
- Do NOT summarize.
- Do NOT answer questions.
- Do NOT solve questions.
- Do NOT explain anything.
- Do NOT add information.
- Do NOT remove information.
- Translate all readable natural-language text.
- Preserve question numbers.
- Preserve headings.
- Preserve names.
- Preserve dates.
- Preserve numbers.
- Preserve IDs.
- Preserve choices such as (a), (b), (c), (d).
- Preserve formulas and mathematical expressions.
- Preserve the order of the document.
- Preserve line breaks where reasonably possible.
- If the source contains more than one language, translate all readable natural-language languages into the target language.
- Do not translate proper names unless they naturally have an established translated form.
- Do not answer questions contained in the document.
- Do not add the word "Translation".
- Return ONLY the translated document text.

DOCUMENT TEXT:

{clean_text}
"""


    try:

        response = client.chat.completions.create(

            model=MODEL,

            messages=[

                {
                    "role": "system",
                    "content":
                        "You are a professional document translator. "
                        "Return only the translation. "
                        "Never summarize, explain, solve, or answer "
                        "questions from the document."
                },

                {
                    "role": "user",
                    "content": prompt
                }

            ],

            temperature=0.1,

            max_tokens=8000

        )


        # -------------------------------------------------
        # Safely get response content
        # -------------------------------------------------

        if not response:

            print(
                "Groq returned no response."
            )

            return ""


        choice_count = len(
            response.choices
        )

        print(
            "Groq choices:",
            choice_count
        )


        if choice_count == 0:

            print(
                "Groq returned zero choices."
            )

            return ""


        message = response.choices[0].message


        if message is None:

            print(
                "Groq message is None."
            )

            return ""


        result = message.content


        if result is None:

            print(
                "Groq message.content is None."
            )

            return ""


        result = str(
            result
        ).strip()


        print(
            "Translation result length:",
            len(result)
        )

        print(
            "Translation preview:",
            result[:300]
        )


        return result


    except Exception as e:

        print(
            "GROQ TRANSLATION ERROR:"
        )

        print(
            str(e)
        )

        return ""


# =========================================================
# LARGE TEXT TRANSLATION
# =========================================================

def translate_large_text(
    text,
    target_language,
    source_language="Auto Detect"
):

    if not text or not text.strip():

        return "No readable text was found."


    # -----------------------------------------------------
    # Smaller chunks are safer for OCR documents
    # -----------------------------------------------------

    chunk_size = 4000


    chunks = []


    for start in range(
        0,
        len(text),
        chunk_size
    ):

        chunk = text[
            start:
            start + chunk_size
        ].strip()


        if chunk:

            chunks.append(
                chunk
            )


    if not chunks:

        return "No readable text was found."


    print("=" * 60)

    print(
        "TOTAL TRANSLATION CHUNKS:",
        len(chunks)
    )

    print("=" * 60)


    translated_chunks = []


    # =====================================================
    # TRANSLATE EACH CHUNK
    # =====================================================

    for index, chunk in enumerate(
        chunks,
        start=1
    ):

        print(
            f"Translating chunk {index}/{len(chunks)}..."
        )


        translated = translate_text(

            chunk,

            target_language,

            source_language

        )


        # -------------------------------------------------
        # If Groq returns empty result
        # -------------------------------------------------

        if not translated:

            print(
                f"WARNING: Chunk {index} returned empty."
            )


            # Retry once

            retry_prompt = f"""
Translate this text into {target_language}.

Return ONLY the translation.

Do not explain.
Do not summarize.
Do not answer questions.

Text:

{chunk}
"""


            try:

                retry_response = client.chat.completions.create(

                    model=MODEL,

                    messages=[

                        {
                            "role": "system",
                            "content":
                                "Translate only. "
                                "Return only translated text."
                        },

                        {
                            "role": "user",
                            "content": retry_prompt
                        }

                    ],

                    temperature=0.1,

                    max_tokens=8000

                )


                if (
                    retry_response
                    and retry_response.choices
                ):

                    retry_content = (
                        retry_response
                        .choices[0]
                        .message
                        .content
                    )


                    if retry_content:

                        translated = (
                            retry_content
                            .strip()
                        )


            except Exception as e:

                print(
                    "Retry translation error:",
                    str(e)
                )


        # -------------------------------------------------
        # Final fallback
        # -------------------------------------------------

        if not translated:

            print(
                f"Chunk {index} could not be translated."
            )

            translated = (
                "[Translation unavailable for this section]\n"
                + chunk
            )


        translated_chunks.append(
            translated
        )


    return "\n\n".join(
        translated_chunks
    )


# =========================================================
# PDF EXTRACTION
# =========================================================

def extract_pdf_text(
    file_bytes,
    source_language="Auto Detect"
):

    pdf = fitz.open(

        stream=file_bytes,

        filetype="pdf"

    )


    extracted_pages = []

    ocr_languages_used = []


    try:

        for page_number, page in enumerate(
            pdf,
            start=1
        ):

            print("=" * 70)

            print(
                f"PROCESSING PDF PAGE {page_number}"
            )

            print("=" * 70)


            # =================================================
            # NORMAL PDF TEXT
            # =================================================

            normal_text = page.get_text(
                "text"
            )


            if (
                normal_text
                and normal_text.strip()
            ):

                print(
                    f"Page {page_number}: Normal PDF text found."
                )


                extracted_pages.append(

                    f"--- Page {page_number} ---\n"
                    f"{normal_text.strip()}"

                )

                continue


            # =================================================
            # OCR
            # =================================================

            print(
                f"Page {page_number}: Scanned/image PDF."
            )

            print(
                "Starting OCR..."
            )


            try:

                pix = page.get_pixmap(

                    matrix=fitz.Matrix(
                        3,
                        3
                    ),

                    alpha=False

                )


                image_bytes = pix.tobytes(
                    "png"
                )


                image = Image.open(

                    io.BytesIO(
                        image_bytes
                    )

                )


                image = image.convert(
                    "RGB"
                )


                ocr_text, ocr_language = ocr_image(

                    image,

                    source_language

                )


                if ocr_language not in ocr_languages_used:

                    ocr_languages_used.append(
                        ocr_language
                    )


                if ocr_text and ocr_text.strip():

                    extracted_pages.append(

                        f"--- Page {page_number} ---\n"
                        f"{ocr_text.strip()}"

                    )

                else:

                    extracted_pages.append(

                        f"--- Page {page_number} ---\n"
                        "[No readable text found]"

                    )


            except Exception as e:

                print(
                    f"OCR ERROR PAGE {page_number}:",
                    str(e)
                )


                extracted_pages.append(

                    f"--- Page {page_number} ---\n"
                    f"[OCR error: {e}]"

                )


    finally:

        pdf.close()


    extracted_text = "\n\n".join(
        extracted_pages
    ).strip()


    print("=" * 70)

    print(
        "PDF EXTRACTION COMPLETED"
    )

    print(
        "Extracted text length:",
        len(extracted_text)
    )

    print("=" * 70)


    return (

        extracted_text,

        ocr_languages_used

    )


# =========================================================
# DOCX EXTRACTION
# =========================================================

def extract_docx_text(
    file_bytes
):

    document = Document(

        io.BytesIO(
            file_bytes
        )

    )


    parts = []


    for paragraph in document.paragraphs:

        if paragraph.text.strip():

            parts.append(
                paragraph.text
            )


    for table in document.tables:

        for row in table.rows:

            cells = []


            for cell in row.cells:

                cells.append(
                    cell.text.strip()
                )


            if any(cells):

                parts.append(

                    " | ".join(
                        cells
                    )

                )


    return "\n".join(
        parts
    ).strip()


# =========================================================
# OLD DOC EXTRACTION
# =========================================================

def extract_doc_text(
    file_bytes,
    original_filename
):

    temp_path = None

    word = None

    document = None


    try:

        with tempfile.NamedTemporaryFile(

            delete=False,

            suffix=".doc"

        ) as temp_file:

            temp_file.write(
                file_bytes
            )

            temp_path = temp_file.name


        import win32com.client


        word = win32com.client.DispatchEx(
            "Word.Application"
        )


        word.Visible = False


        document = word.Documents.Open(

            os.path.abspath(
                temp_path
            )

        )


        text = document.Content.Text


        return text.strip()


    except ImportError:

        raise ValueError(

            "Legacy .DOC support requires pywin32.\n\n"
            "Install it using:\n"
            "pip install pywin32"

        )


    except Exception as e:

        raise ValueError(

            "Could not read the .DOC file.\n\n"
            "Make sure Microsoft Word is installed.\n\n"
            f"Details: {e}"

        )


    finally:

        try:

            if document is not None:

                document.Close(False)

        except Exception:

            pass


        try:

            if word is not None:

                word.Quit()

        except Exception:

            pass


        try:

            if (
                temp_path
                and os.path.exists(
                    temp_path
                )
            ):

                os.remove(
                    temp_path
                )

        except Exception:

            pass


# =========================================================
# CSV READER
# =========================================================

def read_csv_file(
    file_bytes
):

    try:

        return pd.read_csv(

            io.BytesIO(
                file_bytes
            ),

            dtype=str,

            keep_default_na=False

        )

    except Exception:

        return pd.read_csv(

            io.BytesIO(
                file_bytes
            ),

            encoding="latin-1",

            dtype=str,

            keep_default_na=False

        )


# =========================================================
# CSV VALUE CHECK
# =========================================================

def should_translate_csv_value(
    value
):

    if value is None:

        return False


    value = str(
        value
    )


    if not value.strip():

        return False


    stripped = value.strip()


    if re.fullmatch(

        r"[\d\s.,%+\-()/]+",

        stripped

    ):

        return False


    if re.fullmatch(

        r"[^@\s]+@[^@\s]+\.[^@\s]+",

        stripped

    ):

        return False


    if re.match(

        r"^(https?://|www\.)",

        stripped,

        re.IGNORECASE

    ):

        return False


    return True


# =========================================================
# CSV TRANSLATION
# =========================================================

def translate_csv(
    df,
    target_language,
    source_language="Auto Detect"
):

    result = df.copy()

    translation_cache = {}


    def translate_cell(value):

        value = str(
            value
        )


        if not should_translate_csv_value(
            value
        ):

            return value


        if value in translation_cache:

            return translation_cache[
                value
            ]


        translated = translate_text(

            value,

            target_language,

            source_language

        )


        if not translated:

            translated = value


        translation_cache[
            value
        ] = translated


        return translated


    # =====================================================
    # HEADERS
    # =====================================================

    translated_headers = []


    for header in df.columns:

        translated_headers.append(

            translate_cell(
                header
            )

        )


    result.columns = translated_headers


    # =====================================================
    # CELLS
    # =====================================================

    for row_index in df.index:

        for column in df.columns:

            value = df.at[
                row_index,
                column
            ]


            result.at[
                row_index,
                column
            ] = translate_cell(
                value
            )


    return result


# =========================================================
# HOME
# =========================================================

@app.route("/")
def home():

    installed_languages = (
        get_installed_tesseract_languages()
    )


    return render_template(

        "index.html",

        languages=LANGUAGES,

        installed_tesseract_languages=
            installed_languages

    )


# =========================================================
# TEXT TRANSLATION
# =========================================================

@app.route(
    "/translate",
    methods=["POST"]
)
def translate():

    text = request.form.get(
        "text",
        ""
    ).strip()


    source_language = request.form.get(
        "from_language",
        "English"
    )


    target_language = request.form.get(
        "to_language",
        "Tamil"
    )


    if not text:

        return render_template(

            "index.html",

            languages=LANGUAGES,

            error="Please enter some text.",

            text=text,

            from_language=source_language,

            to_language=target_language

        )


    try:

        translated = translate_large_text(

            text,

            target_language,

            source_language

        )


        return render_template(

            "index.html",

            languages=LANGUAGES,

            translation=translated,

            text=text,

            from_language=source_language,

            to_language=target_language

        )


    except Exception as e:

        print(
            "TEXT TRANSLATION ERROR:",
            str(e)
        )


        return render_template(

            "index.html",

            languages=LANGUAGES,

            text=text,

            from_language=source_language,

            to_language=target_language,

            error=
                f"Translation error: {str(e)}"

        )


# =========================================================
# FILE TRANSLATION
# =========================================================

@app.route(
    "/file-translate",
    methods=["POST"]
)
def file_translate():

    uploaded_file = request.files.get(
        "file"
    )


    source_language = request.form.get(
        "file_source_language",
        "Auto Detect"
    )


    target_language = request.form.get(
        "file_language",
        "Tamil"
    )


    # =====================================================
    # FILE CHECK
    # =====================================================

    if (
        not uploaded_file
        or uploaded_file.filename == ""
    ):

        return render_template(

            "index.html",

            languages=LANGUAGES,

            error="Please upload a file."

        )


    filename = secure_filename(
        uploaded_file.filename
    )


    extension = os.path.splitext(
        filename
    )[1].lower()


    print("=" * 70)

    print(
        "FILE TRANSLATION STARTED"
    )

    print(
        "File:",
        filename
    )

    print(
        "Extension:",
        extension
    )

    print(
        "Selected Source:",
        source_language
    )

    print(
        "Target:",
        target_language
    )

    print("=" * 70)


    try:

        file_bytes = uploaded_file.read()


        if not file_bytes:

            return render_template(

                "index.html",

                languages=LANGUAGES,

                error="The uploaded file is empty."

            )


        extracted_text = ""

        translated_text = ""

        csv_html = None

        file_type = ""

        ocr_languages = []


        # =================================================
        # PDF
        # =================================================

        if extension == ".pdf":

            print(
                "PROCESSING PDF..."
            )


            (
                extracted_text,
                ocr_languages
            ) = extract_pdf_text(

                file_bytes,

                source_language

            )


            file_type = "PDF"


            # -------------------------------------------------
            # Determine translation source
            # -------------------------------------------------

            translation_source = detect_translation_source(

                extracted_text,

                source_language

            )


            print(
                "Translation source:",
                translation_source
            )


            # -------------------------------------------------
            # IMPORTANT:
            #
            # Translate the extracted OCR text.
            # -------------------------------------------------

            translated_text = translate_large_text(

                extracted_text,

                target_language,

                translation_source

            )


        # =================================================
        # DOCX
        # =================================================

        elif extension == ".docx":

            print(
                "PROCESSING DOCX..."
            )


            extracted_text = extract_docx_text(
                file_bytes
            )


            file_type = "DOCX"


            translation_source = detect_translation_source(

                extracted_text,

                source_language

            )


            translated_text = translate_large_text(

                extracted_text,

                target_language,

                translation_source

            )


        # =================================================
        # DOC
        # =================================================

        elif extension == ".doc":

            print(
                "PROCESSING LEGACY DOC..."
            )


            extracted_text = extract_doc_text(

                file_bytes,

                filename

            )


            file_type = "DOC"


            translation_source = detect_translation_source(

                extracted_text,

                source_language

            )


            translated_text = translate_large_text(

                extracted_text,

                target_language,

                translation_source

            )


        # =================================================
        # TXT
        # =================================================

        elif extension == ".txt":

            print(
                "PROCESSING TXT..."
            )


            try:

                extracted_text = file_bytes.decode(
                    "utf-8"
                )


            except UnicodeDecodeError:

                extracted_text = file_bytes.decode(
                    "latin-1"
                )


            file_type = "TXT"


            translation_source = detect_translation_source(

                extracted_text,

                source_language

            )


            translated_text = translate_large_text(

                extracted_text,

                target_language,

                translation_source

            )


        # =================================================
        # CSV
        # =================================================

        elif extension == ".csv":

            print(
                "PROCESSING CSV..."
            )


            df = read_csv_file(
                file_bytes
            )


            file_type = "CSV"


            translated_df = translate_csv(

                df,

                target_language,

                source_language

            )


            translated_text = translated_df.to_csv(
                index=False
            )


            csv_html = translated_df.to_html(

                index=False,

                classes="translated-table",

                border=0

            )


        # =================================================
        # UNSUPPORTED
        # =================================================

        else:

            return render_template(

                "index.html",

                languages=LANGUAGES,

                error=(

                    "Unsupported file type.\n\n"

                    "Supported formats:\n"

                    "PDF\n"
                    "DOC\n"
                    "DOCX\n"
                    "TXT\n"
                    "CSV"

                )

            )


        # =================================================
        # EXTRACTED TEXT CHECK
        # =================================================

        if not extracted_text.strip():

            extracted_text = (
                "No readable text was found."
            )


        # =================================================
        # TRANSLATION CHECK
        # =================================================

        if not translated_text.strip():

            translated_text = (

                "Translation could not be produced. "
                "Please check the terminal/console for "
                "the Groq API error."

            )


        # =================================================
        # OCR INFORMATION
        # =================================================

        if ocr_languages:

            ocr_info = (

                "OCR language used: "

                + ", ".join(
                    ocr_languages
                )

            )

        else:

            ocr_info = (
                "OCR was not required."
            )


        print("=" * 70)

        print(
            "TRANSLATION COMPLETED"
        )

        print(
            ocr_info
        )

        print(
            "Extracted characters:",
            len(extracted_text)
        )

        print(
            "Translated characters:",
            len(translated_text)
        )

        print("=" * 70)


        # =================================================
        # OUTPUT FOLDER
        # =================================================

        output_folder = (
            "translated_files"
        )


        os.makedirs(

            output_folder,

            exist_ok=True

        )


        # =================================================
        # FILE NAME
        # =================================================

        base_name = os.path.splitext(
            filename
        )[0]


        # =================================================
        # CSV OUTPUT
        # =================================================

        if extension == ".csv":

            download_filename = (

                f"{base_name}_translated.csv"

            )


            download_path = os.path.join(

                output_folder,

                download_filename

            )


            translated_df.to_csv(

                download_path,

                index=False,

                encoding="utf-8-sig"

            )


        # =================================================
        # OTHER FILES
        # =================================================

        else:

            download_filename = (

                f"{base_name}_translated.txt"

            )


            download_path = os.path.join(

                output_folder,

                download_filename

            )


            with open(

                download_path,

                "w",

                encoding="utf-8"

            ) as output_file:

                output_file.write(
                    translated_text
                )


        print(
            "Download file:",
            download_filename
        )


        # =================================================
        # RETURN RESULT
        # =================================================

        return render_template(

            "index.html",

            languages=LANGUAGES,

            file_name=filename,

            file_type=file_type,

            file_source_language=
                source_language,

            file_language=
                target_language,

            file_source_text=
                extracted_text,

            file_translation=
                translated_text,

            csv_html=
                csv_html,

            download_file=
                download_filename,

            ocr_info=
                ocr_info,

            installed_tesseract_languages=
                get_installed_tesseract_languages()

        )


    except Exception as e:

        print("=" * 70)

        print(
            "FILE PROCESSING ERROR"
        )

        print(
            str(e)
        )

        print("=" * 70)


        return render_template(

            "index.html",

            languages=LANGUAGES,

            error=(

                f"File processing error:\n\n"
                f"{str(e)}"

            ),

            installed_tesseract_languages=
                get_installed_tesseract_languages()

        )


# =========================================================
# DOWNLOAD
# =========================================================

@app.route(
    "/download/<filename>"
)
def download_file(
    filename
):

    return send_from_directory(

        "translated_files",

        filename,

        as_attachment=True

    )


# =========================================================
# RUN APPLICATION
# =========================================================

if __name__ == "__main__":

    app.run(

        host="0.0.0.0",

        port=5000,

        debug=True

    )
